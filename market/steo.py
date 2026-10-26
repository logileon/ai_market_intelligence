"""Versioned EIA STEO archives. No API key; originals are retained byte-for-byte."""
from datetime import datetime, timezone
from hashlib import sha256
from html.parser import HTMLParser
from io import BytesIO
import math
from pathlib import Path
import re
from urllib.parse import urljoin, urlsplit
from zipfile import ZipFile
import pandas as pd
from openpyxl import load_workbook
from .collect import get_bytes
from .storage import ROOT, DB, connect, initialize, log_run

ARCHIVE_URL = "https://www.eia.gov/outlooks/steo/outlook.php"
MONTHS = dict(zip(["jan","feb","mar","apr","may","jun","jul","aug","sep","oct","nov","dec"], range(1,13)))
# Explicit sheet/series mapping avoids ambiguous repeated "World total" labels.
SPECS = {
    "papr_world": ("3atab", "World liquid fuels supply", "million b/d", "flow"),
    "patc_world": ("3atab", "World liquid fuels demand", "million b/d", "flow"),
    "papr_opec": ("3atab", "OPEC liquid fuels supply", "million b/d", "flow"),
    "papr_nonopec": ("3atab", "Non-OPEC liquid fuels supply", "million b/d", "flow"),
    "copr_world": ("3atab", "World crude oil production", "million b/d", "flow"),
    "world_nc": ("3atab", "Other liquid fuels", "million b/d", "flow"),
    "t3_stchange_world": ("3atab", "EIA stock draw & balancing items (+ = draw)", "million b/d", "flow"),
    "patc_oecd": ("3atab", "OECD demand", "million b/d", "flow"),
    "patc_non_oecd": ("3atab", "Non-OECD demand", "million b/d", "flow"),
    "patc_ch": ("3atab", "China demand", "million b/d", "flow"),
    "patc_in": ("3etab", "India demand", "million b/d", "flow"),
    "patc_us": ("3atab", "US liquid fuels demand", "million b/d", "flow"),
    "patc_r03": ("3etab", "Europe demand", "million b/d", "flow"),
    "patc_rs": ("3etab", "Russia demand", "million b/d", "flow"),
    "WTIPUUS": ("2tab", "WTI — EIA forecast", "USD/bbl", "price"),
    "BREPUUS": ("2tab", "Brent — EIA forecast", "USD/bbl", "price"),
    "MGWHUUS_$": ("2tab", "US wholesale gasoline — EIA", "USD/gal", "price"),
    "DSWHUUS_$": ("2tab", "US wholesale diesel — EIA", "USD/gal", "price"),
    "JKTCUUS_$": ("2tab", "US wholesale jet fuel — EIA", "USD/gal", "price"),
    "MGRARUS_$": ("2tab", "US retail regular gasoline — EIA", "USD/gal", "price"),
    "DSRTUUS_$": ("2tab", "US retail on-highway diesel — EIA", "USD/gal", "price"),
    "MGTCPUSX": ("4atab", "US motor gasoline demand", "million b/d", "flow"),
    "DFTCPUS": ("4atab", "US distillate fuel demand", "million b/d", "flow"),
    "JFTCPUS": ("4atab", "US jet fuel demand", "million b/d", "flow"),
}
DERIVED = {"balance": ("Supply minus demand", "million b/d", "flow")}
CORE = {"papr_world", "patc_world", "WTIPUUS", "BREPUUS"}

class ArchiveParser(HTMLParser):
    def __init__(self):
        super().__init__(); self.rows = []; self.text = []; self.links = []; self.in_row = False
    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self.in_row = True; self.text = []; self.links = []
        if tag == "a" and self.in_row:
            href = dict(attrs).get("href", "")
            if re.search(r"/?[a-z]{3}\d{2}_base\.xlsx$", href, re.I): self.links.append(href)
    def handle_data(self, data):
        if self.in_row: self.text.append(data)
    def handle_endtag(self, tag):
        if tag == "tr" and self.in_row:
            self.rows.append((" ".join(self.text), self.links)); self.in_row = False

def discover_releases(html, as_of=None):
    parser = ArchiveParser(); parser.feed(html)
    today = as_of or datetime.now(timezone.utc).date()
    releases = {}
    for text, links in parser.rows:
        if not links: continue
        date_match = re.search(r"\b(\d{1,2})/(\d{1,2})/(\d{4})\b", text)
        if not date_match: continue
        released = datetime(int(date_match[3]), int(date_match[1]), int(date_match[2])).date()
        if released > today: continue
        for href in links:
            url = urljoin(ARCHIVE_URL, href)
            p = urlsplit(url)
            if p.scheme != "https" or p.hostname != "www.eia.gov" or not p.path.startswith("/outlooks/steo/archives/"): continue
            match = re.search(r"/([a-z]{3})(\d{2})_base\.xlsx$", p.path, re.I)
            if not match or match[1].lower() not in MONTHS: continue
            month = f"{2000+int(match[2]):04d}-{MONTHS[match[1].lower()]:02d}"
            if released.strftime("%Y-%m") != month: raise ValueError(f"Archive month mismatch: {url}")
            releases[month] = dict(release_month=month, released_at=released.isoformat(), source_url=url)
    if not releases: raise ValueError("EIA archive list contains no dated XLSX releases")
    return sorted(releases.values(), key=lambda x: x["release_month"], reverse=True)

def initialize_steo(path=DB):
    initialize(path)
    with connect(path) as con:
        con.executescript("""
        CREATE TABLE IF NOT EXISTS steo_vintages (
          vintage_id TEXT PRIMARY KEY, release_month TEXT NOT NULL, released_at TEXT NOT NULL,
          completed_at TEXT NOT NULL, historical_cutoff TEXT NOT NULL, retrieved_at TEXT NOT NULL,
          source_url TEXT NOT NULL, sha256 TEXT NOT NULL, raw_path TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS steo_values (
          vintage_id TEXT NOT NULL, series TEXT NOT NULL, period TEXT NOT NULL, value REAL NOT NULL,
          is_forecast INTEGER NOT NULL, source_sheet TEXT NOT NULL, source_row INTEGER NOT NULL,
          source_cell TEXT NOT NULL, PRIMARY KEY(vintage_id,series,period),
          FOREIGN KEY(vintage_id) REFERENCES steo_vintages(vintage_id));
        """)

def month_key(value):
    try:
        n = int(value)
        year, month = divmod(n, 100)
        return datetime(year, month, 1).strftime("%Y-%m")
    except (ValueError, TypeError, OverflowError):
        raise ValueError(f"Invalid STEO month: {value!r}")

def parse_workbook(body, release):
    with ZipFile(BytesIO(body)) as zipped:
        if sum(i.file_size for i in zipped.infolist()) > 100_000_000:
            raise ValueError("Uncompressed workbook exceeds 100 MB")
    book = load_workbook(BytesIO(body), read_only=True, data_only=True)
    try:
        dates = list(book["Dates"].values)
        metadata = {str(row[0]).strip(): row for row in dates if row[0]}
        def date_field(prefix):
            found = [r[3] for k, r in metadata.items() if k.startswith(prefix)]
            if len(found) != 1: raise ValueError(f"Missing Dates field: {prefix}")
            return found[0]
        stated = str(date_field("Forecast Month")).split()
        expected = datetime.strptime(" ".join(stated), "%B %Y").strftime("%Y-%m")
        if expected != release["release_month"]: raise ValueError("Workbook vintage differs from archive link")
        completed = date_field("Modeling and analysis completion")
        if not isinstance(completed, datetime): raise ValueError("Missing modeling completion date")
        cutoff = month_key(date_field("Last Historical Month"))
        headers = [r for r in dates if len(r)>1 and r[1] == "AAAA_DATEX or AAAA_YEAR"]
        if len(headers) != 1: raise ValueError("Missing monthly date axis")
        columns = {i: month_key(v) for i,v in enumerate(headers[0]) if i >= 2 and v is not None}
        if len(columns) < 24 or len(set(columns.values())) != len(columns):
            raise ValueError("Invalid monthly date axis")
        digest = sha256(body).hexdigest()
        vintage_id = release["release_month"] + "-" + digest[:12]
        observations = []
        found = set()
        for sheet in sorted({v[0] for v in SPECS.values()}):
            if sheet not in book.sheetnames: raise ValueError(f"Missing source sheet: {sheet}")
            seen = {}
            for row in book[sheet].iter_rows():
                code = str(row[0].value or "").strip()
                if code not in SPECS or SPECS[code][0] != sheet: continue
                values = []
                for i, period in columns.items():
                    if i >= len(row): raise ValueError(f"Column missing in {sheet}")
                    cell = row[i]; val = cell.value
                    if val is None or (isinstance(val, str) and val.strip() in ("", "-", "--", "NA", "N/A")): continue
                    if not isinstance(val, (int, float)) or not math.isfinite(val):
                        raise ValueError(f"Non-numeric value: {sheet}!{cell.coordinate}")
                    values.append((vintage_id, code, period, float(val), int(period > cutoff), sheet, row[0].row, cell.coordinate))
                if code in seen:
                    if [(v[2],v[3]) for v in seen[code]] != [(v[2],v[3]) for v in values]:
                        raise ValueError(f"Conflicting duplicate series {code}")
                    continue
                if not values: raise ValueError(f"Empty series {code}")
                seen[code] = values; found.add(code); observations.extend(values)
        if found != set(SPECS): raise ValueError(f"Missing STEO series: {sorted(set(SPECS)-found)}")
        frame = pd.DataFrame(observations, columns=["vintage_id","series","period","value","is_forecast","source_sheet","source_row","source_cell"])
        wide = frame.pivot(index="period", columns="series", values="value")
        # Source's stock-change row is a DRAW sign convention, opposite our supply-demand balance.
        check = (wide.papr_world - wide.patc_world + wide.t3_stchange_world).dropna()
        if check.empty or check.abs().max() > 0.02: raise ValueError("STEO supply/demand/stock-draw reconciliation failed")
        if any(frame[frame.series == code].is_forecast.sum() < 1 for code in CORE):
            raise ValueError("Core series has no forecast periods")
        meta = dict(release, vintage_id=vintage_id, completed_at=completed.date().isoformat(),
                    historical_cutoff=cutoff, retrieved_at=datetime.now(timezone.utc).isoformat(),
                    sha256=digest, raw_path=f"steo/raw/{vintage_id}.xlsx")
        return meta, observations
    finally:
        book.close()

def ingest(body, release, path=DB, data_dir=None):
    initialize_steo(path)
    meta, rows = parse_workbook(body, release)
    base = Path(data_dir) if data_dir else Path(path).parent
    raw = base / meta["raw_path"]
    raw.parent.mkdir(parents=True, exist_ok=True)
    if not raw.exists():
        tmp = raw.with_suffix(".tmp")
        tmp.write_bytes(body); tmp.replace(raw)
    elif sha256(raw.read_bytes()).hexdigest() != meta["sha256"]:
        raise ValueError("Stored original checksum differs")
    with connect(path) as con:
        if con.execute("SELECT 1 FROM steo_vintages WHERE vintage_id=?", (meta["vintage_id"],)).fetchone(): return meta["vintage_id"], 0
        con.execute("""INSERT INTO steo_vintages VALUES
          (:vintage_id,:release_month,:released_at,:completed_at,:historical_cutoff,:retrieved_at,:source_url,:sha256,:raw_path)""", meta)
        con.executemany("INSERT INTO steo_values VALUES (?,?,?,?,?,?,?,?)", rows)
    return meta["vintage_id"], len(rows)

def read_steo(path=DB):
    initialize_steo(path)
    with connect(path) as con:
        return (pd.read_sql_query("SELECT * FROM steo_vintages ORDER BY release_month DESC, retrieved_at DESC, vintage_id DESC",con),
                pd.read_sql_query("SELECT * FROM steo_values ORDER BY vintage_id,series,period",con))

def update_steo(path=DB, history=7, refresh_latest=True):
    initialize_steo(path)
    now = datetime.now(timezone.utc).isoformat()
    results = []
    try:
        releases = discover_releases(get_bytes(ARCHIVE_URL).decode("utf-8", errors="replace"))[:history]
        known, _ = read_steo(path)
        for index, release in enumerate(releases):
            month = release["release_month"]
            try:
                cached = known[known.release_month == month]
                # Raw originals and normalized rows must both exist before treating a version as cached.
                intact = (not cached.empty and all(
                    (Path(path).parent/r.raw_path).exists() and sha256((Path(path).parent/r.raw_path).read_bytes()).hexdigest() == r.sha256
                    for r in cached.itertuples()))
                if intact and not (index == 0 and refresh_latest):
                    status, items, detail = "ok", 0, "Cached immutable vintage; no download"
                else:
                    blob = get_bytes(release["source_url"])
                    vid, items = ingest(blob, release, path)
                    status, detail = "ok", f"{vid}; added observations: {items}"
            except Exception as exc:
                status, items, detail = "error", 0, f"{type(exc).__name__}: {exc}"
            result = dict(source=f"EIA STEO {month}", status=status, items=items, detail=detail)
            log_run(now, **result, path=path); results.append(result)
    except Exception as exc:
        result = dict(source="EIA STEO archive", status="error", items=0, detail=f"{type(exc).__name__}: {exc}")
        log_run(now, **result, path=path); results.append(result)
    return results

def series_meta(code):
    return (SPECS[code][1:] if code in SPECS else DERIVED[code])

def aggregate(values, frequency="Month"):
    """Full-period averages only: flows day-weighted, prices arithmetic monthly means."""
    if values.empty:
        return pd.DataFrame(columns=["vintage_id","series","period","value","status","forecast_months","months"])
    frame = values.copy()
    frame["date"] = pd.to_datetime(frame.period + "-01")
    frame["days"] = frame.date.dt.days_in_month
    frame["group"] = frame.date.dt.to_period({"Month":"M","Quarter":"Q","Year":"Y"}[frequency]).astype(str)
    expected = {"Month":1,"Quarter":3,"Year":12}[frequency]
    output = []
    for (vintage, code, period), group in frame.groupby(["vintage_id","series","group"]):
        complete = len(group) == expected and group.period.nunique() == expected
        value = None
        if complete:
            value = (float((group.value*group.days).sum()/group.days.sum())
                     if series_meta(code)[2] == "flow" else float(group.value.mean()))
        n_forecast = int(group.is_forecast.sum())
        status = "EIA forecast" if n_forecast == expected else "History / estimate" if n_forecast == 0 else "History + forecast"
        if not complete: status = "Incomplete period"
        output.append(dict(vintage_id=vintage, series=code, period=period, value=value, status=status,
                           forecast_months=n_forecast, months=len(group)))
    result = pd.DataFrame(output)
    result["value"] = pd.to_numeric(result["value"], errors="coerce")
    # Balance uses aligned aggregated supply/demand only; no fill or extrapolation.
    p = result[result.series.isin(["papr_world","patc_world"])].pivot(index=["vintage_id","period"],columns="series",values="value")
    if {"papr_world","patc_world"}.issubset(p.columns):
        balance = (p.papr_world-p.patc_world).rename("value").reset_index()
        flags = result[result.series == "papr_world"][["vintage_id","period","status","forecast_months","months"]]
        balance = balance.merge(flags, on=["vintage_id","period"],validate="one_to_one")
        balance["series"] = "balance"
        result = pd.concat([result,balance],ignore_index=True)
    return result.sort_values(["vintage_id","series","period"]).reset_index(drop=True)

def compare(aggregated, current, baseline, series):
    """Outer alignment: missing coverage remains NaN; delta = current minus baseline."""
    left = aggregated[(aggregated.vintage_id == current)&(aggregated.series == series)][["period","value","status"]]
    right = aggregated[(aggregated.vintage_id == baseline)&(aggregated.series == series)][["period","value","status"]]
    out = left.merge(right,on="period",how="outer",suffixes=("_current","_baseline"),validate="one_to_one").sort_values("period")
    out["delta"] = out.value_current - out.value_baseline
    out["delta_pct"] = out.delta.div(out.value_baseline.where(out.value_baseline != 0)).mul(100)
    return out

