"""English edition regression checks, including persisted and exported content."""
from pathlib import Path
from io import BytesIO
import re
import unittest
from openpyxl import load_workbook
from market.storage import ROOT, connect, read_tables
from market.export import excel_bytes
from market.newsapi import parse_articles

CYRILLIC = re.compile(r'[\u0400-\u04ff]')

class EnglishEditionTests(unittest.TestCase):
    def test_interface_and_docs_have_no_cyrillic(self):
        files=[ROOT/'app.py',*list((ROOT/'market').glob('*.py')),*list((ROOT/'config').glob('*.json')),*list(ROOT.glob('*.md'))]
        for file in files:
            with self.subTest(file=file.name):
                self.assertFalse(CYRILLIC.search(file.read_text(encoding='utf-8-sig')))

    def test_saved_news_and_excel_are_english(self):
        news,prices,runs=read_tables()
        self.assertGreater(len(news),0)
        for frame in [news,prices,runs]:
            for row in frame.itertuples(index=False,name=None):
                self.assertTrue(all(not isinstance(v,str) or not CYRILLIC.search(v) for v in row))
        book=load_workbook(BytesIO(excel_bytes(news,prices,runs)),read_only=True)
        for sheet in book:
            for row in sheet.iter_rows(values_only=True):
                self.assertTrue(all(not isinstance(v,str) or not CYRILLIC.search(v) for v in row))
        book.close()

    def test_english_collector_omits_cyrillic_articles(self):
        item=dict(title='Oil '+chr(0x43d)+chr(0x435)+chr(0x444)+chr(0x442)+chr(0x44c),url='https://example.org/a',description='crude oil')
        self.assertEqual(parse_articles(dict(status='ok',articles=[item]),'2026-10-24T12:00:00+00:00'),[])

if __name__=='__main__':unittest.main()
