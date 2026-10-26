from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
import sqlite3
import tempfile
import unittest
from unittest.mock import patch
import pandas as pd
from openpyxl import load_workbook
from market.analysis import analyze, relevant
from market.collect import canonical_url, clean, parse_feed, parse_prices, update
from market.export import excel_bytes
from market.storage import initialize, save_news, read_tables, connect

NOW = datetime.now(timezone.utc).isoformat()
RSS = b'''<rss version="2.0"><channel><title>Test</title>
<item><title>Refinery outage threatens diesel supply</title><link>https://example.org/story?utm_source=rss</link>
<pubDate>Mon, 05 Jan 2026 10:00:00 GMT</pubDate><description>&lt;p&gt;US crude refinery shutdown&lt;/p&gt;</description></item>
<item><title>Natural gas prices</title><link>https://example.org/gas</link></item>
<item><title>Crude oil expansion</title><link>javascript:alert(1)</link></item>
</channel></rss>'''

class PipelineTests(unittest.TestCase):
    def test_relevance_and_evidence(self):
        self.assertTrue(relevant("OPEC crude oil outlook"))
        self.assertFalse(relevant("natural gas and LNG market"))
        result = analyze("Diesel refinery outage in China")
        self.assertIn("Diesel & gasoil", result["topics"])
        self.assertEqual(result["signal"], "Risk mention")
        self.assertIn("outage", result["evidence"])
        self.assertEqual(result["regions"], "Asia")

    def test_rss_and_deduplication(self):
        rows = parse_feed(RSS, "Fixture", NOW)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["url"], "https://example.org/story")
        self.assertEqual(clean("<script>bad</script><p>Oil &amp; products</p>"), "Oil & products")
        self.assertEqual(canonical_url("javascript:alert(1)"), "")
        with tempfile.TemporaryDirectory() as d:
            db = Path(d)/"test.db"; initialize(db)
            self.assertEqual(save_news(rows, db), 1)
            self.assertEqual(save_news(rows, db), 0)
            self.assertEqual(len(read_tables(db)[0]), 1)

    def test_price_missing_and_multi_series(self):
        a = parse_prices(b"observation_date,DCOILWTICO\n2026-01-01,71\n2026-01-02,.\n", "DCOILWTICO")
        b = parse_prices(b"observation_date,DGASNYH\n2026-01-01,2.4\n", "DGASNYH")
        self.assertEqual(len(a), 1)
        with tempfile.TemporaryDirectory() as d:
            db = Path(d)/"test.db"; initialize(db)
            with connect(db) as con:
                con.executemany("INSERT OR REPLACE INTO prices VALUES (?,?,?,?)", a+b)
            self.assertEqual(len(read_tables(db)[1]), 2)
        with self.assertRaises(ValueError): parse_prices(b"hello,world\n1,2", "DCOILWTICO")

    def test_partial_outage_preserves_history(self):
        with tempfile.TemporaryDirectory() as d:
            db = Path(d)/"test.db"; initialize(db)
            save_news(parse_feed(RSS, "Fixture", NOW), db)
            with patch("market.collect.get_bytes", side_effect=ConnectionError("offline")):
                result = update(db)
            self.assertTrue(all(x["status"] in ("error", "skipped") for x in result))
            news, _, runs = read_tables(db)
            self.assertEqual(len(news), 1)
            self.assertEqual(len(runs), 10)

    def test_excel_injection_and_roundtrip(self):
        frame = pd.DataFrame({"title": ["=HYPERLINK(\"bad\")", " +1", "Normal"]})
        blob = excel_bytes(frame, pd.DataFrame(), pd.DataFrame())
        book = load_workbook(BytesIO(blob))
        self.assertEqual(book.sheetnames, ["News", "Prices", "Collection status", "Crack spreads", "Read me"])
        self.assertEqual(book["News"]["A2"].data_type, "s")
        self.assertTrue(book["News"]["A2"].value.startswith("'="))

if __name__ == "__main__": unittest.main()
